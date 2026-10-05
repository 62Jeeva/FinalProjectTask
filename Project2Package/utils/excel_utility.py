from openpyxl import load_workbook


class ExcelUtility:

    def __init__(self,file_path):
        self.file_path = file_path

    def get_sheet(self,sheet_name):
        workbook = load_workbook(self.file_path)
        return workbook[sheet_name]

    def get_all_data(self,sheet_name):
        workbook = load_workbook(self.file_path)
        sheet =workbook[sheet_name]
        headers = [cell.value
                   for cell in sheet[1]
                   ]
        data = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
          row_data = dict(zip(headers,row))
          data.append(row_data)

        return data

    def get_row_data(self,sheet_name,row_number):
        workbook = load_workbook(self.file_path)
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        values = [cell.value for cell in sheet[row_number]]
        return dict(zip(headers,values))

    