from openpyxl import load_workbook

def read_login_data(file_path,sheet_name):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    login_data = []
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        data = (
            row[0].value,
            row[1].value,
            row[2].value,
            row[3].value
        )
        login_data.append(data)

    workbook.close()

    return login_data